"""Génère des CEP d'exemple (données fictives) à partir du modèle V11-20.

Usage : python scripts/make_demo.py <modele.xlsm> <dossier_sortie>

Les valeurs sont écrites directement dans les cellules lues par l'application
(le classeur n'est pas recalculé) : ces fichiers servent uniquement à la démonstration.
"""
import datetime as dt
import random
import sys
from pathlib import Path

import openpyxl

REF = dt.datetime(2026, 9, 1)
PROJECTS = [
    # code, client, projet, DP, CP, typo, facturation, TJM, PRJ, CA mensuel, marge visée, début, fin
    ("DEMO-ACME", "ACME (exemple)", "Refonte ERP", "JDU", "MLE", "GOLD", "Forfait", 720, 520, 600_000, 0.29, (2025, 1), (2027, 6)),
    ("DEMO-BETA", "Beta Assurances (exemple)", "TMA applicative", "SBO", "KTR", "SILVER", "Régie", 640, 500, 420_000, 0.22, (2024, 3), (2026, 12)),
    ("DEMO-GAMMA", "Gamma Énergie (exemple)", "Migration cloud", "JDU", "PNA", "GOLD", "Echéancier", 780, 540, 370_000, 0.31, (2025, 9), (2027, 3)),
    ("DEMO-DELTA", "Delta Retail (exemple)", "Plateforme data", "AMR", "LGI", "BRONZE", "Forfait", 690, 560, 260_000, 0.12, (2026, 1), (2026, 12)),
    ("DEMO-EPSI", "Epsilon Santé (exemple)", "Infogérance", "SBO", "FDA", "SILVER", "Régie", 610, 470, 200_000, 0.24, (2023, 6), (2027, 5)),
    ("DEMO-ZETA", "Zeta Transports (exemple)", "Appli mobile", "AMR", "OCL", "BRONZE", "Forfait", 700, 530, 135_000, 0.06, (2026, 2), (2026, 11)),
]
TEAM = [("Lyon", "MLE", "MARTIN, Léa", "Chef de projet"), ("Lyon", "KTR", "KEROUANI, Tarik", "Développeur senior"),
        ("Rennes", "PNA", "PONS, Nadia", "Architecte"), ("Lyon", "LGI", "GARNIER, Luc", "Data engineer"),
        ("Rennes", "FDA", "DALMAS, Fanny", "Ingénieur run"), ("Lyon", "OCL", "CLEMENT, Omar", "Développeur"),
        ("Lyon", "RBE", "BENALI, Rania", "Testeur")]
SHAPE = [.78, .84, .95, .92, .8, 1, .97, .62, 1, 1.02, 1, .9]


def month_cols(ws, row):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if isinstance(v, dt.datetime) and v.year >= 2000:
            out[(v.year, v.month)] = c
    return out


def report_col(ws, row):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.strip().lower().startswith("report"):
            return c
    return None


def active(ym, start, end):
    return start <= ym <= end


def build(template, out_dir, spec, rnd):
    code, client, projet, dp, cp, typo, fact, tjm, prj, base, mg, start, end = spec
    wb = openpyxl.load_workbook(template, data_only=True, keep_vba=True)
    som, cep, act = wb["Sommaire"], wb["CEP"], wb["Activité"]

    cols = month_cols(cep, 1)
    months = sorted(cols)
    rep = report_col(cep, 1)
    ref = (REF.year, REF.month)

    rows = {k: {} for k in ["sal", "ach", "loc", "pf", "ste", "sti", "tch", "cap", "rev", "cse", "csi", "cab", "rfa", "can", "mb", "mn"]}
    for ym in months:
        if not active(ym, start, end):
            for k in rows:
                rows[k][ym] = 0
            continue
        s = SHAPE[ym[1] - 1] * (0.93 + rnd.random() * 0.14)
        ca = base * s
        cap, rev = ca * 0.86, ca * 0.10
        cse, csi = ca * 0.04, 0
        cout = ca * (1 - mg + (rnd.random() - 0.5) * 0.08)
        sal, ach, ste = cout * 0.74, cout * 0.10, cout * 0.16
        loc = ach * 0.3
        rows["sal"][ym], rows["ach"][ym], rows["loc"][ym], rows["pf"][ym] = sal, ach, loc, 0
        rows["ste"][ym], rows["sti"][ym], rows["tch"][ym] = ste, 0, sal + ach + ste
        rows["cap"][ym], rows["rev"][ym], rows["cse"][ym], rows["csi"][ym] = cap, rev, cse, csi
        rows["cab"][ym] = cap + rev + cse + csi
        rows["rfa"][ym] = rows["cab"][ym] * 0.02
        rows["can"][ym] = rows["cab"][ym] - rows["rfa"][ym]
        rows["mb"][ym] = rows["cab"][ym] - rows["tch"][ym]
        rows["mn"][ym] = rows["can"][ym] - rows["tch"][ym]

    # Report avant 2026 : 12 mois « moyens » par année écoulée avant le démarrage du tableau.
    years_before = max(0, 2026 - start[0]) - (start[1] - 1) / 12
    report = {k: sum(v.values()) / max(1, len(months)) * 12 * years_before for k, v in rows.items()}

    line = {"sal": 85, "ach": 86, "loc": 87, "pf": 89, "ste": 91, "sti": 92, "tch": 93, "cap": 98, "rev": 99, "cse": 100,
            "csi": 101, "cab": 102, "rfa": 111, "can": 112, "mb": 130, "mn": 141}
    for k, r in line.items():
        for ym, c in cols.items():
            cep.cell(r, c).value = round(rows[k][ym], 2)
        if rep:
            cep.cell(r, rep).value = round(report[k], 2)
        cep.cell(r, 6).value = round(sum(rows[k].values()) + report[k], 2)

    # Activité : réalisé Pilot jusqu'au mois de référence, prévisionnel sur toute la période.
    acols = month_cols(act, 1)
    arep = report_col(act, 1)
    members = rnd.sample(TEAM, 4)
    day_tot = {}
    for i, (site, trig, nom, role) in enumerate(members):
        share = [0.34, 0.28, 0.22, 0.16][i]
        for r0, real_only in ((9 + i, True), (46 + i, False)):
            act.cell(r0, 2).value, act.cell(r0, 5).value, act.cell(r0, 6).value, act.cell(r0, 8).value = site, trig, nom, role
            for ym, c in acols.items():
                if not active(ym, start, end) or (real_only and ym > ref):
                    act.cell(r0, c).value = 0
                    continue
                days = rows["cap"].get(ym, 0) / tjm * share * (1 if real_only else 0.92 + rnd.random() * 0.12)
                act.cell(r0, c).value = round(days, 1)
                if real_only:
                    day_tot[ym] = day_tot.get(ym, 0) + days
            if arep:
                act.cell(r0, arep).value = round(report["cap"] / tjm * share, 1) if real_only else 0
    for ym, c in acols.items():
        act.cell(94, c).value = round(day_tot.get(ym, 0) if ym <= ref else rows["cap"].get(ym, 0) / tjm, 1)

    total_days = sum(rows["cap"].values()) / tjm + report["cap"] / tjm
    conso_cumul = sum(v for ym, v in day_tot.items()) + report["cap"] / tjm
    jv = round(total_days * 1.02, -1)
    obj = round(total_days * (0.97 if mg > 0.15 else 0.92), -1)
    raf = max(0, total_days - conso_cumul)
    derive = (conso_cumul + raf - obj)

    def tot(k, upto=None):
        return report[k] + sum(v for ym, v in rows[k].items() if upto is None or ym <= upto)

    renta = lambda upto=None: (tot("mb", upto) / tot("cab", upto)) if tot("cab", upto) else 0
    renta_n = lambda upto=None: (tot("mn", upto) / tot("can", upto)) if tot("can", upto) else 0
    rm = rows["mb"].get(ref, 0) / rows["cab"].get(ref, 1) if rows["cab"].get(ref) else 0

    s = som
    s["C6"], s["C7"], s["C8"], s["C9"] = code, client, projet, "Centre de Delivery"
    s["C10"], s["C11"], s["C12"], s["C13"], s["C14"] = f"AFF-{rnd.randint(10000, 99999)}", "RCO", dp, cp, typo
    s["H7"], s["H8"], s["H9"] = "Services", fact, "Site OPEN"
    s["C19"], s["H19"] = dt.datetime(start[0], start[1], 1), dt.datetime(end[0], end[1], 28)
    s["C20"], s["C21"] = dt.datetime(2026, 10, 1 + rnd.randint(0, 14), 9, 30), REF
    s["C23"] = f"{projet} pour {client.replace(' (exemple)', '')} — données fictives de démonstration."
    s["F28"], s["G28"], s["H28"] = tjm * 1.01, tjm, tjm * 0.99
    s["F29"], s["G29"], s["H29"] = prj, prj * 1.01, prj * 1.02
    s["F30"], s["G30"], s["H30"] = rm, renta(ref), renta()
    s["F33"], s["G33"], s["H33"] = rm - 0.02, renta_n(ref), renta_n()
    s["F31"], s["G31"], s["H31"] = rm + 0.03, renta(ref) + 0.03, renta() + 0.03
    s["C37"], s["H36"] = jv, obj
    s["F37"], s["G37"], s["H37"] = day_tot.get(ref, 0), conso_cumul, conso_cumul + raf
    s["C38"], s["C39"] = jv - (conso_cumul + raf), (jv - (conso_cumul + raf)) / jv
    s["H38"], s["G38"], s["G39"] = derive, derive / obj, raf
    s["H40"], s["H41"], s["H42"] = tot("cab"), tot("can"), tot("tch")
    s["G43"] = 0.55 + rnd.random() * 0.35
    wb["Paramètres"]["B5"] = 0.25

    bc = wb["Bons de commande"]
    for i, (lib, part) in enumerate([("Lot 1 · cadrage", 0.2), ("Lot 2 · réalisation", 0.65), ("Avenant 1", 0.15)]):
        r = 3 + i
        bc.cell(r, 6).value = f"4500{rnd.randint(100000, 999999)}"
        bc.cell(r, 8).value = lib
        bc.cell(r, 9).value = round(jv * part)
        bc.cell(r, 10).value = round(jv * part * tjm, 2)
        bc.cell(r, 11).value = 0
        bc.cell(r, 12).value = ["Oui", "Oui", "Non"][i]
        bc.cell(r, 13).value = tjm
        bc.cell(r, 19).value = tjm * 1.05
        bc.cell(r, 20).value = round(obj * part)

    fac = wb["Facturation"]
    fcols = month_cols(fac, 1)
    for ym, c in fcols.items():
        th = rows["cab"].get(ym, 0) * 0.95
        fac.cell(72, c).value = round(th, 2)
        fac.cell(73, c).value = round(th * (0.9 + rnd.random() * 0.15), 2) if ym <= ref else 0

    pro = wb["Projection"]
    pcols = month_cols(pro, 7)
    for ym, c in pcols.items():
        pro.cell(10, c).value = round(rows["cap"].get(ym, 0) / tjm, 1)
        pro.cell(12, c).value = round(rows["can"].get(ym, 0) * 1.03, 2)
        pro.cell(15, c).value = round(rows["tch"].get(ym, 0) * 0.98, 2)

    pam = wb["PAM"]
    pam["E9"], pam["F9"], pam["G9"], pam["I9"], pam["J9"], pam["N9"], pam["O9"] = (
        "Automatisation des tests de non-régression", cp, "Fort", 42, "En cours", 30, "Oui")
    if mg < 0.15:
        pam["E22"], pam["F22"], pam["G22"], pam["H22"], pam["I22"], pam["J22"], pam["O22"] = (
            "Négociation des écarts au contrat avec le client", dp, "Moyen", 60, 60, "Proposé", "Non")

    target = out_dir / f"{code}_CEP_000_CompteExploitationProjet_FR_V11-20.xlsm"
    wb.save(target)
    return target


def main():
    template, out_dir = Path(sys.argv[1]), Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    rnd = random.Random(7)
    for spec in PROJECTS:
        print("écrit", build(template, out_dir, spec, rnd))


if __name__ == "__main__":
    main()
